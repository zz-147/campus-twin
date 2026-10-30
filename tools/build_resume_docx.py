# -*- coding: utf-8 -*-
"""生成中文简历 DOCX（AI 全栈 / AI Agent / 大模型应用方向）。

用法（DSH 自带 Python，含 python-docx）：
  <bundled python> tools/build_resume_docx.py
输出：docs/简历-校园网络数字孪生运维平台.docx
之后用 LibreOffice Kit 导出 PDF。
"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "简历-校园网络数字孪生运维平台.docx"

DARK = RGBColor(0x1F, 0x2A, 0x3A)
ACCENT = RGBColor(0x0B, 0x5C, 0xA8)
GRAY = RGBColor(0x59, 0x63, 0x70)
FONT = "微软雅黑"


def set_cn_font(run, size=None, bold=None, color=None, name=FONT):
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def para(doc, text="", size=9.5, bold=False, color=DARK, align=WD_ALIGN_PARAGRAPH.LEFT,
         space_before=0, space_after=2, line=1.2, indent=None):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line
    if indent is not None:
        pf.left_indent = Cm(indent)
    for i, seg in enumerate(text.split("\n")):
        if i:
            p.add_run().add_break()
        run = p.add_run(seg)
        set_cn_font(run, size=size, bold=bold, color=color)
    return p


def section_title(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_cn_font(run, size=11.5, bold=True, color=ACCENT)
    # 底部细线
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "0B5CA8")
    pBdr.append(bottom)
    pPr.append(pBdr)


def bullet(doc, text, size=9.5):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.45)
    pf.first_line_indent = Cm(-0.35)
    pf.space_before = Pt(0)
    pf.space_after = Pt(1.5)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.2
    run = p.add_run("• " + text)
    set_cn_font(run, size=size, color=DARK)
    return p


def labeled(doc, label, text, size=9.5):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.45)
    pf.first_line_indent = Cm(-0.45)
    pf.space_before = Pt(0)
    pf.space_after = Pt(1.5)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.2
    r1 = p.add_run("▪ " + label + "：")
    set_cn_font(r1, size=size, bold=True, color=DARK)
    r2 = p.add_run(text)
    set_cn_font(r2, size=size, color=DARK)
    return p


doc = Document()
# A4 页面 + 窄边距
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(1.5)
sec.bottom_margin = Cm(1.4)
sec.left_margin = Cm(1.8)
sec.right_margin = Cm(1.8)

# ---------------- 顶部：姓名 + 求职意向 + 联系方式 ----------------
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
run = p.add_run("【你的姓名】")
set_cn_font(run, size=24, bold=True, color=DARK)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
run = p.add_run("求职意向：AI Agent 开发 / 大模型应用开发 / AI 全栈    |    工作年限与岗位方向按实填写")
set_cn_font(run, size=9.5, bold=True, color=ACCENT)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
run = p.add_run("电话：【手机号】  ·  邮箱：【邮箱】  ·  GitHub：github.com/zz-147  ·  城市：【城市】")
set_cn_font(run, size=9, color=GRAY)

# ---------------- 专业技能 ----------------
section_title(doc, "专业技能")
labeled(doc, "AI Agent / 大模型应用",
        "LangGraph（StateGraph 条件路由 / Send 并行 / interrupt 人机中断 / SqliteSaver checkpointer）、LangChain Core、"
        "function calling 与 tool use、多 Agent 编排、SSE 流式输出、轻量 RAG（检索增强 + [K1] 引用）、Prompt 工程、Agent 评测（38 场景）")
labeled(doc, "模型与协议",
        "DeepSeek / OpenAI 兼容 API；provider 可插拔（DeepSeek / OpenAI / Ollama / mock）；"
        "MCP（Model Context Protocol）：自建 MCP Server + MCP Client 桥接；token 成本控制与超时降级")
labeled(doc, "后端与工程化",
        "Python / FastAPI / asyncio / SQLite WAL / WebSocket；pytest（后端 171 + Agent 256 + MCP 40）；"
        "软删除 + 审计 + 动作确认 / TTL / CAS 等变更安全设计")
labeled(doc, "前端与可视化",
        "Three.js / Vite / 原生 JavaScript；3D 数字孪生大屏、路径追踪动画、SSE 逐 token 渲染、RAG 引用交互、action 确认卡片")
labeled(doc, "网络与数据采集",
        "Cisco Packet Tracer 9 / Network Controller REST API、Cisco IOS-XE（3650 / 2960 / ISR4331）、VLAN / Trunk / OSPF、SNMP 驱动、故障注入实验")

# ---------------- 项目经历 ----------------
section_title(doc, "项目经历")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(1)
r1 = p.add_run("校园网络数字孪生运维平台")
set_cn_font(r1, size=11, bold=True, color=DARK)
r2 = p.add_run("    |    个人项目（独立开发）    |    2026.10    |    Python / LangGraph / MCP / FastAPI / Three.js")
set_cn_font(r2, size=9, color=GRAY)

para(doc, "项目描述：以 Cisco Packet Tracer 9 的 Network Controller REST API 为真实数据源，构建校园网络「采集—检测—归因—"
          "建议—验证」运维闭环；AI 侧用 LangGraph 编排多 Agent，支持流式对话、RAG 引用、人机确认写操作与 MCP 工具协议化，"
          "前端为 Three.js 3D 数字孪生大屏（真实运行，可演示）。", size=9.5, color=GRAY, line=1.22)

bullet(doc, "真实数据采集：5 秒周期采集 5 台网络设备与 6 台终端（拓扑 / 端口 / 可达性 / 健康问题），5 台设备全部 Managed；"
            "设计 7 条检测规则与告警生命周期，故障注入实测 5 秒内产生 critical 告警。")
bullet(doc, "告警治理：针对 NC 告警清除滞后（实测 >3 分钟）改用主机 ping / 设备可达性独立判定恢复并标记 stale_source；"
            "统一 observed_at / source_ts 双时间基准解决仿真时钟偏差；按根因分组（link_down > device_unreachable > "
            "device_offline > host_down）并抑制下游告警，避免告警风暴。")
bullet(doc, "LangGraph 多 Agent 编排：StateGraph 节点 + 条件边实现意图分类（规则优先 + LLM）→ 8 个 role 并行 fan-out → 汇总；"
            "SqliteSaver checkpointer（thread_id=conversation_id）支持多轮续聊与进程重启恢复；"
            "interrupt() / Command(resume) 实现写操作人机确认（危险操作必须人工确认，10 分钟 TTL + CAS 防重 + 审计）。")
bullet(doc, "流式与 RAG：SSE 逐 token 输出（真实 LLM 实测单次 2600+ token 事件、前端 80ms 节流渲染、可停止/自动回退）；"
            "18 个知识块中文 bigram BM25-lite 检索 + [K1] 引用，无命中明确回答「知识库无相关内容」，不编造。")
bullet(doc, "MCP 双向集成：把 10 个只读工具 / 4 个资源 / 3 个提示词封装为 MCP Server（stdio），Claude Desktop / Cursor / "
            "Inspector 任一 MCP 客户端可接入；另以 MCP Client 桥接（独立 Python 3.14 进程 + HTTP 适配层）为 Agent 接入 "
            "filesystem / memory 外部 MCP Server（23 工具），只读白名单 + 调用审计。")
bullet(doc, "工程与质量：后端 pytest 171、Agent 256、MCP 40 全离线可跑；38 场景评测集（LangGraph / 自研双引擎）根因命中 100%、"
            "设备定位 100%、动作正确率 100%、自动执行违规 0；支持 DRIVER=netcon / SNMP 可插拔，为接真实校园设备预留接口。")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(1)
r1 = p.add_run("【其他项目 / 实习经历（建议补充）】")
set_cn_font(r1, size=11, bold=True, color=DARK)
r2 = p.add_run("    |    时间    |    技术栈")
set_cn_font(r2, size=9, color=GRAY)
para(doc, "项目描述：【一句话说明背景 + 你的职责 + 成果】", size=9.5, color=GRAY)
bullet(doc, "【核心技术点 1：问题 → 方案 → 量化结果】")
bullet(doc, "【核心技术点 2：问题 → 方案 → 量化结果】")

# ---------------- 开源项目（GitHub） ----------------
section_title(doc, "开源项目（GitHub: github.com/zz-147）")
bullet(doc, "shuxi-data-agent（数析·智能数据分析平台）：自然语言 → 计划 → 受控工具（SQL / 指标 / 归因 / 图表 / 报告）→ 结论，"
            "带 HITL 审批、注入防护与全链路 Trace（Python + Streamlit + 通义千问）。")
bullet(doc, "smart-vending-platform（智购柜·智能售货机管理平台核心模块）：XXL-Job 分片补货调度 / Redisson 分布式锁 / "
            "库存乐观锁 / RAG 推荐与智能客服（Spring Boot 3.3 + Java 21）。")
bullet(doc, "deepsearch-agents：带来源追溯的一主三从深度研搜多智能体工作流；shopkeeper-agent：安全的电商智能问数，"
            "中文问题到受控 SQL 的可测试工作流。")

# ---------------- 教育经历 ----------------
section_title(doc, "教育经历")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
r1 = p.add_run("【学校名称】    |    【专业】    |    【学历（本科/硕士）】    |    【入学时间 — 毕业时间】")
set_cn_font(r1, size=9.8, bold=True, color=DARK)
para(doc, "主修课程：【数据结构、计算机网络、操作系统、数据库等，可选填】    绩点/排名：【可选填】", size=9.3, color=GRAY)

# ---------------- 荣誉证书 ----------------
section_title(doc, "荣誉与证书")
bullet(doc, "【证书 / 奖项 / 竞赛：如 CCNA、软考、AI 相关证书、竞赛名次，按实填写】")

# ---------------- 自我评价 ----------------
section_title(doc, "自我评价")
bullet(doc, "工程落地能力强：能独立完成「网络环境搭建 → 后端采集 → Agent 编排 → 前端可视化 → 测试评测」全链路，项目代码 400+ 测试全离线可跑。")
bullet(doc, "对 AI Agent 有真实实践：LangGraph 编排、人机中断、checkpointer、MCP 协议化、SSE 流式、RAG 评测均有可演示的实现与数据。")
bullet(doc, "注重数据真实性与可验证性：拒绝「套壳式 Demo」，关键结论都有实测数据与截图支撑；面试可现场演示真实故障注入与 AI 归因。")

# 页脚提示（可选）
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(8)
run = p.add_run("提示：【】内为占位内容，填写后删除本行；本项目细节见 GitHub / 作品集 PDF。")
set_cn_font(run, size=8, color=RGBColor(0x99, 0x99, 0x99))

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print("已生成:", OUT)

# 校验
check = Document(OUT)
texts = [p.text for p in check.paragraphs if p.text.strip()]
print("段落数:", len(texts), "| 字符数:", sum(len(t) for t in texts))
print("前 6 段:")
for t in texts[:6]:
    print("  -", t[:80])
